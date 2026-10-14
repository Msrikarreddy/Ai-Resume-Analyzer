
from io import BytesIO

import pymupdf
from docx import Document


def extract_resume_text(uploaded_file):
    """Extract text from a PDF or DOCX resume."""

    file_bytes = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        with pymupdf.open(
            stream=file_bytes,
            filetype="pdf"
        ) as pdf:
            text = "\n".join(
                page.get_text(sort=True)
                for page in pdf
            )

    elif filename.endswith(".docx"):
        document = Document(BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        # Extract text from tables too.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)

        text = "\n".join(paragraphs)

    else:
        raise ValueError(
            "Please upload a PDF or DOCX file."
        )

    if not text.strip():
        raise ValueError(
            "No readable text found. "
            "The PDF may be scanned."
        )

    return text
    